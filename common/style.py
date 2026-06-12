#!/usr/bin/env python3
"""Procalc shared spreadsheet styling.

Palette + cell helpers + the "live unit label" formula helpers shared by
the PSV and CV tools, so both follow an identical visual theme:

  * white sheet background throughout
  * light-grey bands mark titles, section headers and table column headers
    (subtle — no heavy colour bars)
  * cyan-coloured text on white marks user-editable input/field cells
  * everything else (labels, calculated results, units, notes) is plain
    dark text on white
  * PASS / FAIL results keep their green / red highlight

Importing modules add ``common/`` to ``sys.path`` before ``import style``
(same convention as ``import units as UN``).
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import units as UN

# ═══════════════════════════════════════════════════════════════════════════
#  Palette
# ═══════════════════════════════════════════════════════════════════════════
WHITE  = "FFFFFF"   # sheet background — "everything else"
LGRAY  = "F7F7F7"   # faint tint: label column, inactive-row shading
NAVY   = "F2F2F2"   # title bar / column-header band (light grey, not flashy)
GRNHDR = "F2F2F2"   # section-header band (same light grey)
AMBER  = "FFFFFF"   # input-cell fill — the cyan font marks the field
EGRAY  = "FFFFFF"   # calculated-cell fill
TEAL   = "FFFFFF"   # generic result-value fill
DGRAY  = "1A1A1A"   # primary text colour
ORANGE = "C55A11"   # warning / flag text colour
CYAN   = "0070C0"   # input / field font colour


def _fillc(h):
    return PatternFill("solid", fgColor=h)


def _bord():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def C(ws, r, c, v=None, bg=None, fg=DGRAY, sz=10, bold=False, wrap=False,
      ha="left"):
    x = ws.cell(row=r, column=c)
    if v is not None:
        x.value = v
    x.font = Font(name="Calibri", size=sz, bold=bold, color=fg)
    if bg:
        x.fill = _fillc(bg)
    x.alignment = Alignment(horizontal=ha, vertical="center", wrap_text=wrap)
    x.border = _bord()
    return x


def _title(ws, ncol, text):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=1 + ncol)
    C(ws, 1, 2, text, bg=NAVY, fg=DGRAY, sz=12, bold=True)
    ws.row_dimensions[1].height = 22


# ── live unit labels (track the UNITS sheet) ────────────────────────────────
# UNITS sheet layout (see common/units.py add_units_sheet): "Unit System" in
# B3, then one row per quantity code (column A) — in QUANTITY_NAMES order,
# starting row 6 — with the per-quantity override in column C.
UNIT_ROWS = {qty: 6 + i for i, qty in enumerate(UN.QUANTITY_NAMES.keys())}


class _U:
    """Marks a `_kv_block`/header unit as a *live* formula referencing the
    UNITS sheet, rather than literal text fixed at template-creation time.

    ``_U("P")``            → the current label for quantity "P"
    ``_U("Q", "/", "dT")``  → composite, e.g. "kW/°C" (literals pass through)
    """
    __slots__ = ("parts",)

    def __init__(self, *parts):
        self.parts = parts


def _pretty_unit(u: str) -> str:
    return {"degF": "°F", "degC": "°C"}.get(u, u)


def _unit_expr(qty: str) -> str:
    """Excel expression (no leading '=') for the live label of `qty`:
    UNITS!C<row> override if set, else the system default for UNITS!B3."""
    row = UNIT_ROWS.get(qty)
    if row is None:
        return f'"{_pretty_unit(UN.SYSTEM_DEFAULTS["FPS"].get(qty, ""))}"'
    ovr = f"UNITS!$C${row}"
    fps_def = _pretty_unit(UN.SYSTEM_DEFAULTS["FPS"][qty])
    si_def = _pretty_unit(UN.SYSTEM_DEFAULTS["SI"][qty])
    pretty_ovr = f'IF({ovr}="degF","°F",IF({ovr}="degC","°C",{ovr}))'
    default = f'IF(UNITS!$B$3="SI","{si_def}","{fps_def}")'
    return f'IF({ovr}<>"",{pretty_ovr},{default})'


def _unit_formula(qty: str) -> str:
    return "=" + _unit_expr(qty)


def _unit_formula_parts(parts) -> str:
    pieces = [_unit_expr(p) if p in UNIT_ROWS else f'"{p}"' for p in parts]
    return "=" + " & ".join(pieces)


def _unit_cell(unit):
    """Resolve a `_kv_block` unit-hint: `_U(...)` → live formula string,
    anything else passes through unchanged (literal text)."""
    if isinstance(unit, _U):
        if len(unit.parts) == 1:
            return _unit_formula(unit.parts[0])
        return _unit_formula_parts(unit.parts)
    return unit


def _hdr_formula(base: str, qty: str) -> str:
    """Live column-header formula: '<base> (<live unit>)'."""
    return f'="{base} (" & {_unit_expr(qty)} & ")"'


def _kv_block(ws, r, rows, widths=(34, 16, 12, 64)):
    """rows = list of (label, default, unit_hint, note, kind) — kind: in/calc/sec."""
    for col, w in zip("BCDE", widths):
        ws.column_dimensions[col].width = w
    for label, default, unit, note, kind in rows:
        if kind == "sec":
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
            C(ws, r, 2, label, bg=GRNHDR, fg=DGRAY, sz=10, bold=True)
        else:
            C(ws, r, 2, label, bg=LGRAY, sz=9, bold=False)
            C(ws, r, 3, default,
              bg=(EGRAY if kind == "calc" else AMBER),
              fg=(DGRAY if kind == "calc" else CYAN), sz=9, ha="right")
            C(ws, r, 4, _unit_cell(unit), sz=9, fg="808080")
            C(ws, r, 5, note, sz=9, fg="808080", wrap=True)
        r += 1
    return r


def _result_rows(ws, r0, rows, widths=(34, 30, 10, 56)):
    """Render a key/value result block.

    ``rows`` = list of ``(label, value, unit, note)``; ``note == "sec"``
    marks a section-header row.  ``value`` of literal ``"PASS"``/``"FAIL"``
    is colour-coded as a compliance result.
    """
    for col, w in zip("BCDE", widths):
        ws.column_dimensions[col].width = w
    for label, val, unit, note in rows:
        if note == "sec":
            ws.merge_cells(start_row=r0, start_column=2, end_row=r0, end_column=5)
            C(ws, r0, 2, label, bg=GRNHDR, fg=DGRAY, sz=10, bold=True)
        else:
            pf = val in ("PASS", "FAIL")
            bg = "C6EFCE" if val == "PASS" else "FFC7CE" if val == "FAIL" else TEAL
            C(ws, r0, 2, label, bg=LGRAY, sz=9)
            C(ws, r0, 3, "—" if val is None else val, bg=bg, sz=9, bold=pf,
              ha="center" if pf else "right")
            C(ws, r0, 4, unit, sz=9, fg="808080")
            C(ws, r0, 5, note, sz=9, fg="808080", wrap=True)
        r0 += 1
    return r0
