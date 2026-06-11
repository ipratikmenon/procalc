#!/usr/bin/env python3
"""Procalc shared units layer — single codebase, FPS or SI presentation.

Every Procalc engine computes internally in FPS canonical units (the proven
numerical base):  psia · °F · lb/hr · lb/ft³ · ft · in (bore/NPS) · in² · ft² ·
ft³ · BTU/hr · BTU/lb · ft/s · cP · lb-mol/hr · psi (ΔP).

A ``UNITS`` sheet in each input workbook declares how the *user* enters values
and how results are *presented*.  The :class:`UnitSystem` object converts at
the I/O boundary only — engine internals never change.

Usage
-----
    from units import UnitSystem, add_units_sheet
    usys = UnitSystem.from_workbook(wb)        # reads the UNITS sheet
    p_int = usys.to_internal("P", p_user)      # user → internal (psia)
    p_out = usys.disp("P", p_int)              # internal → display
    usys.label("P")                            # e.g. "kPa"
    usys.hdr("Start P", "P")                   # e.g. "Start P (kPa)"
"""
from __future__ import annotations

# ── conversion constants (exact / NIST) ────────────────────────────────────
LB_PER_KG = 2.2046226218487757
FT_PER_M = 3.2808398950131233
IN_PER_MM = 1.0 / 25.4
PSI_PER_KPA = 0.14503773773020922
PSIA_ATM = 14.695948775513449
KPA_ATM = 101.325
BTU_PER_KJ = 0.9478171203133172

# quantity → {unit: (to_internal(v), from_internal(v))}
_LIN = lambda f: (lambda v: v * f, lambda v: v / f)  # noqa: E731

_CONVERTERS: dict[str, dict[str, tuple]] = {
    # absolute pressure → psia
    "P": {
        "psia": _LIN(1.0),
        "psig": (lambda v: v + PSIA_ATM, lambda v: v - PSIA_ATM),
        "kPa":  _LIN(PSI_PER_KPA),                      # kPa(a)
        "kPag": (lambda v: (v + KPA_ATM) * PSI_PER_KPA,
                 lambda v: v / PSI_PER_KPA - KPA_ATM),
        "MPa":  _LIN(PSI_PER_KPA * 1000.0),
        "bara": _LIN(PSI_PER_KPA * 100.0),
        "barg": (lambda v: v * 100.0 * PSI_PER_KPA + PSIA_ATM,
                 lambda v: (v - PSIA_ATM) / (100.0 * PSI_PER_KPA)),
        "kg/cm2a": _LIN(14.223343307120154),
        "kg/cm2g": (lambda v: v * 14.223343307120154 + PSIA_ATM,
                    lambda v: (v - PSIA_ATM) / 14.223343307120154),
    },
    # differential pressure → psi
    "dP": {
        "psi":  _LIN(1.0),
        "kPa":  _LIN(PSI_PER_KPA),
        "bar":  _LIN(PSI_PER_KPA * 100.0),
        "mbar": _LIN(PSI_PER_KPA / 10.0),
        "kg/cm2": _LIN(14.223343307120154),
    },
    # temperature → °F
    "T": {
        "degF": _LIN(1.0), "°F": _LIN(1.0),
        "degC": (lambda v: v * 1.8 + 32.0, lambda v: (v - 32.0) / 1.8),
        "°C":   (lambda v: v * 1.8 + 32.0, lambda v: (v - 32.0) / 1.8),
        "K":    (lambda v: (v - 273.15) * 1.8 + 32.0,
                 lambda v: (v - 32.0) / 1.8 + 273.15),
    },
    # temperature difference → °F
    "dT": {"degF": _LIN(1.0), "degC": _LIN(1.8), "K": _LIN(1.8)},
    # mass flow → lb/hr
    "mflow": {
        "lb/hr": _LIN(1.0), "lb/h": _LIN(1.0),
        "kg/hr": _LIN(LB_PER_KG), "kg/h": _LIN(LB_PER_KG),
        "kg/s":  _LIN(LB_PER_KG * 3600.0),
        "t/h":   _LIN(LB_PER_KG * 1000.0),
    },
    # molar flow → lb-mol/hr
    "molflow": {
        "lb-mol/hr": _LIN(1.0),
        "kgmol/hr":  _LIN(LB_PER_KG), "kmol/h": _LIN(LB_PER_KG),
    },
    # density → lb/ft³
    "rho": {
        "lb/ft3": _LIN(1.0), "lb/ft³": _LIN(1.0),
        "kg/m3":  _LIN(LB_PER_KG / FT_PER_M**3),
        "kg/m³":  _LIN(LB_PER_KG / FT_PER_M**3),
    },
    # length → ft
    "L": {"ft": _LIN(1.0), "m": _LIN(FT_PER_M), "mm": _LIN(FT_PER_M / 1000.0)},
    # small length / diameter → in   (NPS stays in inches by convention)
    "Lin": {"in": _LIN(1.0), "mm": _LIN(IN_PER_MM)},
    # area → ft²
    "A": {"ft2": _LIN(1.0), "ft²": _LIN(1.0),
          "m2": _LIN(FT_PER_M**2), "m²": _LIN(FT_PER_M**2)},
    # small area (orifice) → in²
    "Ain": {"in2": _LIN(1.0), "in²": _LIN(1.0),
            "cm2": _LIN((IN_PER_MM * 10.0)**2), "cm²": _LIN((IN_PER_MM * 10.0)**2),
            "mm2": _LIN(IN_PER_MM**2), "mm²": _LIN(IN_PER_MM**2)},
    # volume → ft³
    "V": {"ft3": _LIN(1.0), "ft³": _LIN(1.0),
          "m3": _LIN(FT_PER_M**3), "m³": _LIN(FT_PER_M**3),
          "L": _LIN(FT_PER_M**3 / 1000.0)},
    # heat duty → BTU/hr
    "Q": {"BTU/hr": _LIN(1.0),
          "kW": _LIN(BTU_PER_KJ * 3600.0),
          "W":  _LIN(BTU_PER_KJ * 3.6),
          "MW": _LIN(BTU_PER_KJ * 3.6e6),
          "MMBTU/hr": _LIN(1.0e6), "kcal/h": _LIN(3.968320719312103)},
    # specific enthalpy / latent heat → BTU/lb
    "h": {"BTU/lb": _LIN(1.0), "kJ/kg": _LIN(BTU_PER_KJ / LB_PER_KG),
          "kcal/kg": _LIN(1.7988310432962014)},
    # velocity → ft/s
    "v": {"ft/s": _LIN(1.0), "m/s": _LIN(FT_PER_M)},
    # volumetric flow → ft³/hr
    "qvol": {"ft3/hr": _LIN(1.0), "m3/h": _LIN(FT_PER_M**3),
             "gpm": _LIN(8.020833333333334), "m3/hr": _LIN(FT_PER_M**3)},
    # dimensionless / identical in both systems
    "visc": {"cP": _LIN(1.0)},
    "MW":   {"g/mol": _LIN(1.0)},
    "-":    {"-": _LIN(1.0), "—": _LIN(1.0)},
    "pct":  {"%": _LIN(1.0)},
    "st":   {"dyn/cm": _LIN(1.0), "mN/m": _LIN(1.0)},   # numerically equal
}

# per-system default unit per quantity
SYSTEM_DEFAULTS: dict[str, dict[str, str]] = {
    "FPS": {
        "P": "psia", "dP": "psi", "T": "degF", "dT": "degF",
        "mflow": "lb/hr", "molflow": "lb-mol/hr", "rho": "lb/ft3",
        "L": "ft", "Lin": "in", "A": "ft2", "Ain": "in2", "V": "ft3",
        "Q": "BTU/hr", "h": "BTU/lb", "v": "ft/s", "qvol": "ft3/hr",
        "visc": "cP", "MW": "g/mol", "-": "-", "pct": "%", "st": "dyn/cm",
    },
    "SI": {
        "P": "kPa", "dP": "kPa", "T": "degC", "dT": "degC",
        "mflow": "kg/hr", "molflow": "kgmol/hr", "rho": "kg/m3",
        "L": "m", "Lin": "mm", "A": "m2", "Ain": "cm2", "V": "m3",
        "Q": "kW", "h": "kJ/kg", "v": "m/s", "qvol": "m3/h",
        "visc": "cP", "MW": "g/mol", "-": "-", "pct": "%", "st": "dyn/cm",
    },
}

QUANTITY_NAMES = {
    "P": "Pressure (absolute)", "dP": "Pressure drop", "T": "Temperature",
    "dT": "Temperature difference", "mflow": "Mass flow", "molflow": "Molar flow",
    "rho": "Density", "L": "Length / elevation", "Lin": "Bore / small length",
    "A": "Surface area", "Ain": "Orifice area", "V": "Volume",
    "Q": "Heat duty", "h": "Specific enthalpy / latent heat", "v": "Velocity",
    "qvol": "Volumetric flow", "visc": "Viscosity", "MW": "Molecular weight",
}

UNITS_SHEET = "UNITS"


class UnitSystem:
    """Converts between user-facing units and internal FPS canonical units."""

    def __init__(self, system: str = "FPS",
                 overrides: dict[str, str] | None = None):
        system = (system or "FPS").strip().upper()
        if system not in SYSTEM_DEFAULTS:
            system = "FPS"
        self.system = system
        self.sel: dict[str, str] = dict(SYSTEM_DEFAULTS[system])
        for qty, unit in (overrides or {}).items():
            if qty in _CONVERTERS and unit in _CONVERTERS[qty]:
                self.sel[qty] = unit

    # ── conversion ──────────────────────────────────────────────────────
    def to_internal(self, qty: str, value):
        if value is None or qty not in _CONVERTERS:
            return value
        try:
            v = float(value)
        except (TypeError, ValueError):
            return value
        return _CONVERTERS[qty][self.sel[qty]][0](v)

    def from_internal(self, qty: str, value):
        if value is None or qty not in _CONVERTERS:
            return value
        try:
            v = float(value)
        except (TypeError, ValueError):
            return value
        return _CONVERTERS[qty][self.sel[qty]][1](v)

    def disp(self, qty: str, value, nd: int | None = 4):
        """from_internal + optional rounding (None → no rounding)."""
        out = self.from_internal(qty, value)
        if nd is not None and isinstance(out, float):
            out = round(out, nd)
        return out

    # ── labels ──────────────────────────────────────────────────────────
    def label(self, qty: str) -> str:
        u = self.sel.get(qty, "")
        return {"degF": "°F", "degC": "°C"}.get(u, u)

    def hdr(self, base: str, qty: str) -> str:
        """'Start P', 'P'  →  'Start P (psia)' / 'Start P (kPa)'."""
        lbl = self.label(qty)
        return f"{base} ({lbl})" if lbl and lbl != "-" else base

    @property
    def is_fps(self) -> bool:
        return self.system == "FPS"

    # ── workbook I/O ────────────────────────────────────────────────────
    @classmethod
    def from_workbook(cls, wb) -> "UnitSystem":
        """Read the UNITS sheet (graceful: absent sheet → FPS)."""
        if UNITS_SHEET not in wb.sheetnames:
            return cls("FPS")
        ws = wb[UNITS_SHEET]
        system, overrides = "FPS", {}
        for row in ws.iter_rows(min_row=1, max_row=60, max_col=4,
                                values_only=True):
            if not row or row[0] is None:
                continue
            key = str(row[0]).strip()
            if key.lower() in ("unit system", "system"):
                system = str(row[1] or "FPS").strip().upper()
            elif key in _CONVERTERS and row[2]:
                unit = str(row[2]).strip()
                # normalise pretty labels back to converter keys
                unit = {"°F": "degF", "°C": "degC"}.get(unit, unit)
                overrides[key] = unit
        return cls(system, overrides)


def add_units_sheet(wb, system: str = "FPS", position: int = 0):
    """Insert a styled UNITS sheet into ``wb`` (openpyxl Workbook)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    NAVY, WHITE, LGRAY, AMBER = "1F4973", "FFFFFF", "F2F2F2", "FFF2CC"
    if UNITS_SHEET in wb.sheetnames:
        del wb[UNITS_SHEET]
    ws = wb.create_sheet(UNITS_SHEET, position)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "C00000"

    def cell(r, c, v, bg=None, fg="000000", bold=False, sz=10):
        x = ws.cell(row=r, column=c, value=v)
        x.font = Font(name="Calibri", size=sz, bold=bold, color=fg)
        if bg:
            x.fill = PatternFill("solid", fgColor=bg)
        x.alignment = Alignment(horizontal="left", vertical="center")
        return x

    ws.merge_cells("A1:D1")
    cell(1, 1, "UNITS  —  choose FPS or SI; per-quantity overrides optional",
         bg=NAVY, fg=WHITE, bold=True, sz=11)
    ws.row_dimensions[1].height = 22

    cell(3, 1, "Unit System", bg=LGRAY, bold=True)
    c = cell(3, 2, (system or "FPS").upper(), bg=AMBER, bold=True)
    dv = DataValidation(type="list", formula1='"FPS,SI"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B3")
    cell(3, 3, "← FPS or SI.  All inputs read & all results written in these units.")

    cell(5, 1, "Qty code", bg=NAVY, fg=WHITE, bold=True)
    cell(5, 2, "Quantity", bg=NAVY, fg=WHITE, bold=True)
    cell(5, 3, "Unit override (optional)", bg=NAVY, fg=WHITE, bold=True)
    cell(5, 4, "Options", bg=NAVY, fg=WHITE, bold=True)
    r = 6
    sysd = SYSTEM_DEFAULTS.get((system or "FPS").upper(), SYSTEM_DEFAULTS["FPS"])
    for qty, name in QUANTITY_NAMES.items():
        cell(r, 1, qty, bg=LGRAY)
        cell(r, 2, name)
        cell(r, 3, None, bg=AMBER)   # blank = system default
        opts = " | ".join(_CONVERTERS[qty].keys())
        cell(r, 4, f"default: {sysd[qty]}   |   {opts}", fg="808080", sz=9)
        r += 1
    cell(r + 1, 1, "Leave overrides blank to use the system defaults. "
         "Overrides apply per quantity (e.g. P = barg while system = SI).",
         fg="808080", sz=9)
    for col, w in (("A", 10), ("B", 30), ("C", 22), ("D", 70)):
        ws.column_dimensions[col].width = w
    return ws


# ── self-test ───────────────────────────────────────────────────────────────
if __name__ == "__main__":
    si = UnitSystem("SI")
    fps = UnitSystem("FPS")
    assert abs(si.to_internal("P", 101.325) - PSIA_ATM) < 1e-9
    assert abs(si.to_internal("T", 100.0) - 212.0) < 1e-12
    assert abs(si.to_internal("mflow", 1000.0) - 2204.6226218487758) < 1e-9
    assert abs(si.to_internal("rho", 1000.0) - 62.427960576144606) < 1e-9
    assert abs(si.to_internal("L", 1.0) - FT_PER_M) < 1e-12
    assert abs(si.from_internal("P", si.to_internal("P", 543.21)) - 543.21) < 1e-9
    assert fps.to_internal("P", 100.0) == 100.0
    b = UnitSystem("SI", {"P": "barg"})
    assert abs(b.to_internal("P", 1.0) - (100.0 * PSI_PER_KPA + PSIA_ATM)) < 1e-9
    print("units.py self-test OK")
