"""Streams / HMB viewer: a faithful, sheet-by-sheet rendering of the loaded
HMB workbook -- exactly as the source Excel file shows it, not a hand-built
reconstruction -- with a searchable sheet-name list standing in for tabs
(a real HMB workbook can have 400+ sheets, far too many for a literal
scrolling tab bar to stay usable), and a clickable unit dropdown layered
directly onto any sheet shaped like a real Property/Unit table (every
per-stream sheet, and the workbook's own OUTPUT summary sheet, or a
PRO/II "Case N" transposed export) via streams.unit_sheet_model.
UnitAwareSheetModel.

Every sheet renders through the same results/workbook_model.build_sheet_view
the Results tab uses for engine output -- so a per-stream sheet's merged
title row, column widths, fills and fonts come out identical to opening the
file in Excel.

A live HYSYS/PRO-II connection (engine_api.hmb_source_kind() == "hysys" /
"proii_com") has no on-disk workbook to read -- there, the sheet list is
just the connection's stream names, and selecting one resolves that single
stream live (engine_api.resolve_stream_for_snapshot) and renders it through
streams.live_sheet.build_synthetic_stream_sheet, a sheet built in the SAME
shape a real per-stream sheet uses, so it gets the exact same faithful
rendering and unit dropdowns with no separate code path.

A full-fidelity load of a large HMB workbook takes real time (openpyxl's
non-read_only parser materializes every sheet up front; there is no
partial/lazy-but-full-fidelity mode in its public API), so it runs on a
background QThread once per (path, mtime) and is cached for the app
session. Resolving one live stream over COM can also be slow (first
resolve after connect), so that too runs on a background thread.
"""
from __future__ import annotations

import os

from openpyxl import Workbook, load_workbook
from PySide6.QtCore import Qt, QThread, Signal
from PySide6.QtWidgets import (QDialog, QHBoxLayout, QLabel, QLineEdit,
                               QListWidget, QListWidgetItem, QSplitter,
                               QVBoxLayout, QWidget)

import engine_api as api
from results.workbook_model import build_sheet_view
from streams.live_sheet import build_synthetic_stream_sheet
from streams.unit_sheet_model import UnitAwareSheetModel, install_unit_dropdowns

# Non-per-stream sheets, pinned as a small group at the top of the sheet
# list ahead of the (potentially hundreds of) per-stream sheets below.
_META_SHEETS = ("UNITS", "COMPONENTS", "COMP_CONSTANTS", "INPUT", "OUTPUT")
_LIVE_KINDS = ("hysys", "proii_com")


class _WorkbookLoadWorker(QThread):
    finished_ok = Signal(object)   # openpyxl.Workbook
    failed = Signal(str)

    def __init__(self, path, parent=None):
        super().__init__(parent)
        self._path = path

    def run(self):
        try:
            wb = load_workbook(self._path, read_only=False, data_only=True)
        except Exception as e:  # noqa: BLE001
            self.failed.emit(f"{type(e).__name__}: {e}")
            return
        self.finished_ok.emit(wb)


class _StreamResolveWorker(QThread):
    finished_ok = Signal(str, object)   # name, StreamProps
    failed = Signal(str, str)           # name, message

    def __init__(self, hmb_path, case, name, parent=None):
        super().__init__(parent)
        self._hmb_path = hmb_path
        self._case = case
        self._name = name

    def run(self):
        try:
            sp, _feed = api.resolve_stream_for_snapshot(
                self._hmb_path, self._case, self._name)
        except Exception as e:  # noqa: BLE001
            self.failed.emit(self._name, f"{type(e).__name__}: {e}")
            return
        self.finished_ok.emit(self._name, sp)


class StreamsDialog(QDialog):
    def __init__(self, hmb_path: str, case: str = "Case 1", parent=None,
                workbook_cache: dict | None = None):
        super().__init__(parent)
        self.setWindowTitle("Streams — HMB Workbook")
        self.resize(1300, 800)
        self._hmb_path = hmb_path
        self._case = case
        self._kind = api.hmb_source_kind(hmb_path)
        self._live = self._kind in _LIVE_KINDS
        # shared with the caller (typically MainWindow) so re-opening this
        # dialog for the same file/session is instant instead of re-paying
        # the full-workbook parse; a fresh dict if the caller doesn't pass one.
        self._workbook_cache = workbook_cache if workbook_cache is not None else {}
        self._wb = None                 # file-mode: the loaded workbook
        self._live_wb = None            # live-mode: synthetic sheets accumulate here
        self._current_view = None
        self._worker = None
        self._resolve_worker = None
        self._closed = False

        lay = QVBoxLayout(self)
        lay.setContentsMargins(10, 10, 10, 10)
        lay.setSpacing(8)

        self.status_lbl = QLabel("Loading…")
        self.status_lbl.setObjectName("Muted")
        lay.addWidget(self.status_lbl)

        split = QSplitter(Qt.Horizontal)
        left = QWidget()
        llay = QVBoxLayout(left)
        llay.setContentsMargins(0, 0, 0, 0)
        llay.setSpacing(6)
        self.search = QLineEdit()
        self.search.setPlaceholderText("Filter sheets…")
        self.search.textChanged.connect(self._apply_filter)
        llay.addWidget(self.search)
        self.sheet_list = QListWidget()
        self.sheet_list.currentItemChanged.connect(self._on_sheet_selected)
        llay.addWidget(self.sheet_list, 1)
        left.setMinimumWidth(180)
        left.setMaximumWidth(280)
        split.addWidget(left)

        content_host = QWidget()
        self._content_slot = QVBoxLayout(content_host)
        self._content_slot.setContentsMargins(0, 0, 0, 0)
        split.addWidget(content_host)
        split.setStretchFactor(0, 0)
        split.setStretchFactor(1, 1)
        lay.addWidget(split, 1)

        if self._live:
            self._load_live()
        else:
            self._load_file()

    def closeEvent(self, event):
        self._closed = True
        super().closeEvent(event)

    # ── file-backed HMB (per-stream dump / Case-N export) ──────────────────
    def _cache_key(self):
        try:
            return (self._hmb_path, os.path.getmtime(self._hmb_path))
        except OSError:
            return None

    def _load_file(self):
        self.status_lbl.setText("Loading HMB workbook…")
        key = self._cache_key()
        cached = self._workbook_cache.get(key) if key else None
        if cached is not None:
            self._on_loaded(cached)
            return
        self._worker = _WorkbookLoadWorker(self._hmb_path, self)
        self._worker.finished_ok.connect(self._on_loaded)
        self._worker.failed.connect(self._on_load_failed)
        self._worker.start()

    def _on_loaded(self, wb):
        if self._closed:
            return
        self._wb = wb
        key = self._cache_key()
        if key:
            self._workbook_cache[key] = wb
        self.status_lbl.setText(
            f"{len(wb.sheetnames)} sheets · {os.path.basename(self._hmb_path)}")
        self._populate_file_list()

    def _on_load_failed(self, msg):
        if self._closed:
            return
        self.status_lbl.setText(f"Could not load HMB workbook: {msg}")

    def _populate_file_list(self):
        self.sheet_list.clear()
        names = self._wb.sheetnames
        meta = [n for n in names if n in _META_SHEETS]
        streams = [n for n in names if n not in _META_SHEETS]
        for n in meta:
            self.sheet_list.addItem(QListWidgetItem(n))
        if meta and streams:
            sep = QListWidgetItem(f"── {len(streams)} streams ──")
            sep.setFlags(Qt.NoItemFlags)
            self.sheet_list.addItem(sep)
        for n in streams:
            self.sheet_list.addItem(QListWidgetItem(n))

        default_row = 0
        for i in range(self.sheet_list.count()):
            if self.sheet_list.item(i).text() == "OUTPUT":
                default_row = i
                break
        if self.sheet_list.count():
            self.sheet_list.setCurrentRow(default_row)

    def _apply_filter(self, term):
        term = (term or "").strip().lower()
        for i in range(self.sheet_list.count()):
            item = self.sheet_list.item(i)
            item.setHidden(bool(term) and term not in item.text().lower())

    def _on_sheet_selected(self, current, _previous):
        if current is None:
            return
        name = current.text()
        if self._live:
            self._show_live_stream(name)
            return
        if self._wb is not None and name in self._wb.sheetnames:
            self._show_view(build_sheet_view(self._wb[name], model_cls=UnitAwareSheetModel))

    # ── live HYSYS/PRO-II connection ────────────────────────────────────────
    def _load_live(self):
        kind_label = {"hysys": "HYSYS (live)", "proii_com": "PRO/II (live)"}.get(self._kind, "live")
        self.status_lbl.setText(f"Listing streams… ({kind_label})")
        self._live_wb = Workbook()
        self._live_wb.remove(self._live_wb.active)   # drop the default blank sheet
        try:
            names = api.list_streams(self._hmb_path, self._case)
        except Exception as e:  # noqa: BLE001
            self.status_lbl.setText(f"Could not list streams: {type(e).__name__}: {e}")
            return
        self.status_lbl.setText(
            f"{len(names)} streams · {kind_label} · select a stream to resolve it live")
        self.sheet_list.clear()
        for n in names:
            self.sheet_list.addItem(QListWidgetItem(n))
        if self.sheet_list.count():
            self.sheet_list.setCurrentRow(0)

    def _show_live_stream(self, name):
        if name in self._live_wb.sheetnames:
            self._show_view(build_sheet_view(self._live_wb[name], model_cls=UnitAwareSheetModel))
            return
        self.status_lbl.setText(f"Resolving '{name}' live…")
        self._resolve_worker = _StreamResolveWorker(self._hmb_path, self._case, name, self)
        self._resolve_worker.finished_ok.connect(self._on_live_resolved)
        self._resolve_worker.failed.connect(self._on_live_resolve_failed)
        self._resolve_worker.start()

    def _on_live_resolved(self, name, sp):
        if self._closed:
            return
        build_synthetic_stream_sheet(self._live_wb, name, sp)
        self.status_lbl.setText(f"{self.sheet_list.count()} streams · resolved '{name}'")
        current = self.sheet_list.currentItem()
        if current is not None and current.text() == name:
            self._show_view(build_sheet_view(self._live_wb[name], model_cls=UnitAwareSheetModel))

    def _on_live_resolve_failed(self, name, msg):
        if self._closed:
            return
        self.status_lbl.setText(f"Could not resolve '{name}': {msg}")

    # ── sheet content ────────────────────────────────────────────────────
    def _show_view(self, view):
        if self._current_view is not None:
            self._content_slot.removeWidget(self._current_view)
            self._current_view.setParent(None)
            self._current_view.deleteLater()
            self._current_view = None
        install_unit_dropdowns(view, view.model())
        self._content_slot.addWidget(view)
        self._current_view = view
