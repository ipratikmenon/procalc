"""Build an in-memory, single-sheet openpyxl workbook for one live-resolved
stream (HYSYS/PRO-II COM connection) -- there is no on-disk HMB file for a
live connection to render faithfully, so this synthesizes a sheet in the
SAME shape (and the SAME title/header/section-band styling, colors pulled
directly from the user's real per-stream HMB.xlsx sheets) as a real
per-stream HMB sheet: title row, Property | Unit | TOTAL | VAPOR | LIQUID
header, then five numbered, colour-banded sections. This keeps
streams.unit_sheet_model.UnitAwareSheetModel's existing shape-1 detector
working unchanged -- no separate "live" rendering path needed anywhere
else, and the same clickable unit dropdowns work identically -- while
looking visually consistent with the file-backed case instead of a bare
uncoloured table.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import engine_api as api

# Colors match the real per-stream HMB sheet exactly (read directly off a
# real file's cells): title/header band and section "1" share one navy,
# each subsequent section its own color -- the same identifiers
# Hydraulics/hydraulics_XOM.py itself uses for this exact palette
# (NAVY/DGRAY/STEEL/GRNHDR/PURPLE), duplicated here as literals rather
# than importing the engine directly (engine_api is the only importer of
# hydraulics_XOM, per this app's own architecture).
_NAVY, _DGRAY, _STEEL, _GRNHDR, _PURPLE = "1F4973", "404040", "2E75B6", "375623", "7030A0"
_WHITE, _LGRAY = "FFFFFF", "F2F2F2"

_TITLE_FONT = Font(bold=True, color=_WHITE, size=10)
_HEADER_FONT = Font(bold=True, color=_WHITE, size=9)
_SECTION_FONT = Font(bold=True, color=_WHITE, size=9)
_LABEL_FONT = Font(size=9)
_VALUE_FONT = Font(size=9)

_VALUE_COLS = {"TOTAL": 4, "VAPOR": 5, "LIQUID": 6}

# Section title, its band color, and its (label, {column: StreamProps
# field}, quantity code | None) rows -- mirrors the real per-stream
# sheet's own five numbered sections/column convention exactly (a phase
# property lives only in that phase's column, e.g. Vapor Viscosity only
# in the VAPOR column) against whatever fields a live-resolved
# StreamProps actually carries.
_SECTIONS = [
    ("1.  FLOW RATES", _NAVY, [
        ("Mass Rate", {"TOTAL": "total_mass", "VAPOR": "vap_mass", "LIQUID": "liq_mass"}, "mflow"),
        ("Std Liq Vol Rate", {"TOTAL": "total_std_liq"}, "qvol"),
        ("Std Vap Vol Rate", {"TOTAL": "total_std_vap"}, "qvol"),
    ]),
    ("2.  CONDITIONS", _DGRAY, [
        ("Temperature", {"TOTAL": "temp_f"}, "T"),
        ("Pressure", {"TOTAL": "pres_psia"}, "P"),
        ("Molecular Weight", {"TOTAL": "mol_weight", "VAPOR": "vap_mw", "LIQUID": "liq_mw"}, "MW"),
        ("Specific Enthalpy", {"VAPOR": "vap_sp_enthalpy", "LIQUID": "liq_sp_enthalpy"}, "h"),
    ]),
    ("3.  VAPOR PHASE PROPERTIES", _STEEL, [
        ("Actual Density", {"VAPOR": "vap_density"}, "rho"),
        ("Viscosity", {"VAPOR": "vap_visc"}, "visc"),
        ("Cp", {"VAPOR": "vap_cp"}, None),
        ("Cp/Cv Ratio", {"VAPOR": "vap_cp_cv"}, None),
        ("Z Factor (from density)", {"VAPOR": "vap_z"}, None),
        ("Thermal Conductivity", {"VAPOR": "vap_therm_cond"}, None),
    ]),
    ("4.  LIQUID PHASE PROPERTIES", _GRNHDR, [
        ("Actual Density", {"LIQUID": "liq_density"}, "rho"),
        ("Viscosity", {"LIQUID": "liq_visc"}, "visc"),
        ("Cp", {"LIQUID": "liq_cp"}, None),
        ("Surface Tension", {"LIQUID": "liq_surf_tens"}, "st"),
        ("Thermal Conductivity", {"LIQUID": "liq_therm_cond"}, None),
    ]),
    ("5.  CRITICAL PROPERTIES  (used by EOS in hydraulics engine)", _PURPLE, [
        ("True Critical Temperature", {"TOTAL": "tc_f"}, "T"),
        ("True Critical Pressure", {"TOTAL": "pc_psia"}, "P"),
        ("Actual Density (bulk)", {"TOTAL": "total_density"}, "rho"),
        ("Total Z (from actual density)", {"TOTAL": "total_z"}, None),
        ("Acentric Factor", {"TOTAL": "acentric"}, None),
    ]),
]


def build_synthetic_stream_sheet(wb: Workbook, name: str, sp) -> None:
    """Adds one sheet named `name` to `wb`, in the same shape (and colors)
    a real per-stream HMB sheet uses. `sp` is a StreamProps (already
    resolved via engine_api.resolve_stream_for_snapshot) -- every value is
    already in the engine's internal FPS unit, from
    engine_api.internal_unit(qty)."""
    ws = wb.create_sheet(title=str(name)[:31])

    title = f"STREAM: {name}  |  Phase: {sp.phase or '—'}  |  live connection"
    ws.cell(row=1, column=2, value=title)
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=6)
    cell = ws.cell(row=1, column=2)
    cell.fill = PatternFill("solid", fgColor=_NAVY)
    cell.font = _TITLE_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 20

    headers = ["", "Property", "Unit", "TOTAL", "VAPOR", "LIQUID"]
    for c, h in enumerate(headers, start=1):
        if not h:
            continue
        cell = ws.cell(row=2, column=c, value=h)
        cell.fill = PatternFill("solid", fgColor=_NAVY)
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    r = 3
    zebra = 0
    for section_title, color, rows in _SECTIONS:
        ws.cell(row=r, column=2, value=f"  {section_title}")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        sec_cell = ws.cell(row=r, column=2)
        sec_cell.fill = PatternFill("solid", fgColor=color)
        sec_cell.font = _SECTION_FONT
        sec_cell.alignment = Alignment(horizontal="left", vertical="center")
        r += 1

        for label, col_fields, qty in rows:
            bg = _LGRAY if zebra % 2 else _WHITE
            zebra += 1
            for col_name, field in col_fields.items():
                val = getattr(sp, field, None)
                if val is not None:
                    ws.cell(row=r, column=_VALUE_COLS[col_name], value=val)
            lbl_cell = ws.cell(row=r, column=2, value=label)
            lbl_cell.font = _LABEL_FONT
            lbl_cell.fill = PatternFill("solid", fgColor=bg)
            unit = api.raw_hmb_unit_for(qty) if qty else "-"
            unit_cell = ws.cell(row=r, column=3, value=unit or "-")
            unit_cell.font = _LABEL_FONT
            unit_cell.fill = PatternFill("solid", fgColor=bg)
            for col in (4, 5, 6):
                vc = ws.cell(row=r, column=col)
                vc.font = _VALUE_FONT
                vc.fill = PatternFill("solid", fgColor=bg)
                vc.alignment = Alignment(horizontal="right")
            r += 1

    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 10
    for col in "DEF":
        ws.column_dimensions[col].width = 14
