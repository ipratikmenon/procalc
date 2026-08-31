"""Results panel: a circuit selector + a tab per output sheet, each rendered
from the engine-produced .xlsx."""
from __future__ import annotations

import os

from openpyxl import load_workbook
from PySide6.QtCore import Qt
from PySide6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QComboBox,
                               QLabel, QTabWidget)

from results.workbook_model import build_sheet_view
from results.charts import pressure_profile_chart, flow_pattern_chart


def _build_tab(ws, wb):
    """A sheet tab. For chart sheets, pair the table with a QtCharts re-plot in
    an inner Table/Chart QTabWidget. Any chart failure falls back to table-only.
    """
    table = build_sheet_view(ws)
    title = ws.title or ""
    chart_view = None
    try:
        if title.startswith("Pressure_Profile"):
            chart_view = pressure_profile_chart(ws)
        elif title.startswith("H ") or title.startswith("V "):
            data_ws = wb["Flow_Pattern_Data"] if "Flow_Pattern_Data" in wb.sheetnames else None
            chart_view = flow_pattern_chart(ws, data_ws)
    except Exception:
        chart_view = None
    if chart_view is None:
        return table
    inner = QTabWidget()
    inner.setDocumentMode(True)
    inner.addTab(table, "Table")
    inner.addTab(chart_view, "Chart")
    return inner


class ResultsView(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self._paths: list[str] = []
        lay = QVBoxLayout(self)
        lay.setContentsMargins(6, 6, 6, 6)
        top = QHBoxLayout()
        top.addWidget(QLabel("Circuit:"))
        self.circuit_cb = QComboBox()
        # items are added via addItem() *after* the combo box's first show, so
        # the default AdjustToContentsOnFirstShow policy would freeze the box
        # at its near-empty initial width and clip the real item text (e.g.
        # "C1_DXX5-BTM" rendering as "C1_DXX!") — recompute on every change.
        self.circuit_cb.setSizeAdjustPolicy(QComboBox.AdjustToContents)
        self.circuit_cb.setMinimumWidth(160)
        self.circuit_cb.currentIndexChanged.connect(self._show_current)
        top.addWidget(self.circuit_cb)
        top.addStretch(1)
        self.info = QLabel("")
        self.info.setStyleSheet("color:#555;")
        top.addWidget(self.info)
        lay.addLayout(top)
        self.tabs = QTabWidget()
        self.tabs.setDocumentMode(True)
        lay.addWidget(self.tabs, 1)
        self._placeholder()

    def _placeholder(self):
        self.tabs.clear()
        w = QLabel("Run a calculation to see results here.\n"
                   "The output sheets (Cover, Line List, Pressure Profile, "
                   "Flash, CV datasheets, …) appear as tabs.")
        w.setAlignment(Qt.AlignCenter)
        w.setStyleSheet("color:#7a8a99; font-size:13px;")
        self.tabs.addTab(w, "Results")

    def load(self, paths: list[str]):
        self._paths = [p for p in (paths or []) if p and os.path.exists(p)]
        self.circuit_cb.blockSignals(True)
        self.circuit_cb.clear()
        for p in self._paths:
            name = os.path.basename(p).replace("hydraulics_flash_noiso_", "").replace(".xlsx", "")
            self.circuit_cb.addItem(name, p)
        self.circuit_cb.blockSignals(False)
        if self._paths:
            self._show_current()
        else:
            self._placeholder()

    def current_workbook_path(self):
        idx = self.circuit_cb.currentIndex()
        if 0 <= idx < len(self._paths):
            return self.circuit_cb.itemData(idx)
        return None

    def current_sheet_title(self):
        i = self.tabs.currentIndex()
        return self.tabs.tabText(i) if i >= 0 else None

    def _show_current(self):
        idx = self.circuit_cb.currentIndex()
        if idx < 0 or idx >= len(self._paths):
            return
        prev = self.tabs.currentIndex()
        prev_name = self.tabs.tabText(prev) if prev >= 0 else None
        path = self.circuit_cb.itemData(idx)
        self.tabs.clear()
        wb = load_workbook(path, data_only=False)
        for ws in wb.worksheets:
            self.tabs.addTab(_build_tab(ws, wb), ws.title)
        self.info.setText(f"{self.tabs.count()} sheets  ·  {os.path.basename(path)}")
        # restore the previously-open tab by name if present
        if prev_name:
            for i in range(self.tabs.count()):
                if self.tabs.tabText(i) == prev_name:
                    self.tabs.setCurrentIndex(i)
                    break
